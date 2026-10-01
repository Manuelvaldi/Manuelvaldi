#!/usr/bin/env python3
"""Actualiza Perfiles_Explicados con dos cambios pedidos por Daría:

1. Lámina "Características por perfil": la última columna pasa de
   "% en universo Tenpo" a "% de clientes Tenpo" (los 7 perfiles suman 100%).
2. Lámina nueva (después de la 3): gráfico de segmentos SAM. Un grupo por perfil,
   3 columnas por grupo (ingreso financiero neto, ingreso por servicio neto y
   costo de riesgo hacia abajo), ancho del grupo proporcional a las colocaciones
   y, sobre cada grupo, el % de colocaciones de los bancos competidores (BF, BS...).

Los números salen de datos_perfiles.json. Lo que falte se muestra como "pend.".

    python3 build_perfiles.py                       # usa datos_perfiles.json
    python3 build_perfiles.py --demo --out demo.pptx  # datos ILUSTRATIVOS, solo para ver el diseño
"""
import argparse
import copy
import json
import math
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A}

# Paleta del deck (sacada de la lámina 3)
INK, MUTED, ACCENT = "0F172A", "6B7280", "0891B2"
ACCENT_LIGHT, NEG = "7DD3E8", "B42318"
GRID, BAND = "E5E7EB", "F8FAFC"
FONT = "Arial"

PROFILES = [f"P{i}" for i in range(1, 8)]

# Datos ILUSTRATIVOS (solo --demo): sirven para validar el diseño, no son cifras reales.
DEMO = {
    "P1": dict(clientes_tenpo=90_000, colocaciones=4200, ifn=14.0, isn=3.1, costo_riesgo=3.0, competidores=dict(BF=12, BS=8)),
    "P2": dict(clientes_tenpo=210_000, colocaciones=6800, ifn=18.0, isn=4.0, costo_riesgo=5.5, competidores=dict(BF=14, BS=9)),
    "P3": dict(clientes_tenpo=760_000, colocaciones=21000, ifn=31.0, isn=5.2, costo_riesgo=14.0, competidores=dict(BF=22, BS=15)),
    "P4": dict(clientes_tenpo=540_000, colocaciones=18500, ifn=27.0, isn=4.6, costo_riesgo=17.0, competidores=dict(BF=18, BS=14)),
    "P5": dict(clientes_tenpo=140_000, colocaciones=3200, ifn=16.0, isn=3.8, costo_riesgo=2.0, competidores=dict(BF=9, BS=12)),
    "P6": dict(clientes_tenpo=190_000, colocaciones=9800, ifn=12.0, isn=2.4, costo_riesgo=5.0, competidores=dict(BF=10, BS=20)),
    "P7": dict(clientes_tenpo=14_000, colocaciones=600, ifn=9.0, isn=1.9, costo_riesgo=2.5, competidores=dict(BF=6, BS=25)),
}
# En DEMO ifn/isn/costo_riesgo vienen como MM CLP; los armamos como tasa * colocaciones / 100
for _p in DEMO.values():
    for _k in ("ifn", "isn", "costo_riesgo"):
        _p[_k] = _p[_k] / 100 * _p["colocaciones"]


# --------------------------------------------------------------------------- datos
def load_data(path, demo):
    if demo:
        return copy.deepcopy(DEMO)
    raw = json.loads(Path(path).read_text(encoding="utf8"))["perfiles"]
    return {p: raw.get(p, {}) for p in PROFILES}


def pct_clientes(data):
    """% de clientes Tenpo por perfil con 1 decimal y suma exacta 100.0 (método del mayor resto)."""
    counts = {p: data[p].get("clientes_tenpo") for p in PROFILES}
    direct = {p: data[p].get("pct_clientes_tenpo") for p in PROFILES}
    if all(v is not None for v in counts.values()):
        total = sum(counts.values())
        if total <= 0:
            return {p: None for p in PROFILES}
        raw = {p: counts[p] / total * 1000 for p in PROFILES}  # décimas de punto
    elif all(v is not None for v in direct.values()):
        total = sum(direct.values())
        if abs(total - 100) > 0.5:
            sys.exit(f"pct_clientes_tenpo suma {total:.1f}, debe sumar 100")
        raw = {p: direct[p] * 10 for p in PROFILES}
    else:
        return {p: None for p in PROFILES}
    floors = {p: math.floor(v) for p, v in raw.items()}
    missing = 1000 - sum(floors.values())
    for p in sorted(PROFILES, key=lambda q: raw[q] - floors[q], reverse=True)[:missing]:
        floors[p] += 1
    return {p: floors[p] / 10 for p in PROFILES}


def fmt_pct(v, dec=1, sign=False):
    s = f"{v:.{dec}f}%"
    return ("+" + s) if sign and v > 0 else s


# --------------------------------------------------------------------------- helpers pptx
def rgb(h):
    return RGBColor.from_string(h)


def add_rect(slide, x, y, w, h, fill=None, line=None, dash=False, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    # sin referencia de estilo del tema: evita sombras en LibreOffice y herencias raras
    style = s._element.find(f"{{{PML}}}style")
    if style is not None:
        s._element.remove(style)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    return s


def add_text(slide, x, y, w, h, text, size=7, bold=False, color=INK, align="l", anchor="m", vert=None, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if vert:
        tf._txBody.find(f"{{{A}}}bodyPr").set("vert", vert)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        runs = ln if isinstance(ln, list) else [(ln, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = italic
            f.color.rgb = rgb(o.get("color", color))
    return tb


def set_run_text(tc, text, color=None, bold=None):
    """Cambia el texto de la primera corrida de una celda conservando el formato."""
    r = tc.xpath(".//a:r")[0]
    r.find("a:t", NS).text = text
    rpr = r.find("a:rPr", NS)
    if color:
        rpr.find("a:solidFill/a:srgbClr", NS).set("val", color)
    if bold is not None:
        rpr.set("b", "1" if bold else "0")


# --------------------------------------------------------------------------- lámina 3
def update_profile_table(prs, pcts, demo):
    slide = prs.slides[2]
    tbl = next(sh for sh in slide.shapes if sh.has_table).table
    assert tbl.cell(0, 6)._tc.xpath("string(.//a:t)") == "% en universo Tenpo", "la lámina 3 cambió, revisar"
    set_run_text(tbl.cell(0, 6)._tc, "% de clientes Tenpo")
    for i, p in enumerate(PROFILES, start=1):
        tc = tbl.cell(i, 0)._tc
        assert tc.xpath("string(.//a:t)") == p, f"fila {i} no es {p}"
        v = pcts[p]
        if v is None:
            set_run_text(tbl.cell(i, 6)._tc, "pend.", color=MUTED, bold=False)
        else:
            set_run_text(tbl.cell(i, 6)._tc, fmt_pct(v))
    note = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Valores"))
    t = note._element.xpath(".//a:t")[0]
    old = "% en universo Tenpo: personas del perfil de 18-54 años, sin E Sub Bancarizado."
    assert old in t.text, "no encontré la nota al pie esperada"
    t.text = t.text.replace(
        old, "% de clientes Tenpo: distribución de la base de clientes Tenpo entre los 7 perfiles (suma 100%)."
    )
    if demo:
        stamp(slide)
    return slide


def stamp(slide):
    b = add_rect(slide, 7.0, 0.16, 2.6, 0.3, fill="FEE2E2", line=NEG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, 7.0, 0.16, 2.6, 0.3, "DATOS ILUSTRATIVOS · NO USAR", size=8, bold=True, color=NEG, align="c")


# --------------------------------------------------------------------------- lámina SAM
def profile_style(src_slide):
    """Color y nombre de cada perfil, leídos de la tabla de la lámina 3."""
    tbl = next(sh for sh in src_slide.shapes if sh.has_table).table
    out = {}
    for i, p in enumerate(PROFILES, start=1):
        c0 = tbl.cell(i, 0)._tc
        out[p] = dict(
            fill=c0.xpath("string(./a:tcPr/a:solidFill/a:srgbClr/@val)"),
            ink=c0.xpath("string(.//a:rPr/a:solidFill/a:srgbClr/@val)"),
            name=tbl.cell(i, 1)._tc.xpath("string((.//a:t)[1])"),
        )
    return out


def nice_step(span):
    for s in (1, 2, 2.5, 5, 10, 20, 25, 50, 100):
        if span / s <= 6:
            return s
    return 100


def build_sam_slide(prs, data, style, demo):
    src = prs.slides[2]
    slide = prs.slides.add_slide(src.slide_layout)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    # banda + título copiados de la lámina 3 para heredar el estilo exacto
    for sh in list(src.shapes)[:2]:
        slide.shapes._spTree.append(copy.deepcopy(sh._element))
    slide.shapes[1]._element.xpath(".//a:t")[0].text = "Segmentos SAM: ingresos y costo de riesgo por perfil"

    # ---- datos
    need = ("colocaciones", "ifn", "isn", "costo_riesgo")
    complete = all(data[p].get(k) is not None for p in PROFILES for k in need)
    banks = []
    for p in PROFILES:
        for b in (data[p].get("competidores") or {}):
            if b not in banks:
                banks.append(b)
    banks = banks or ["BF", "BS"]

    # ---- geometría (pulgadas)
    X0, X1, GAP = 1.05, 9.59, 0.07
    strip_y = 0.80
    strip_h = 0.16 * len(banks) + 0.06
    plot_top = strip_y + strip_h + 0.10
    plot_bot = 4.40
    chip_y = plot_bot + 0.08
    leg_y = 4.98
    note_y = 5.17

    # anchos proporcionales a colocaciones (con piso para que se lea)
    MIN_W = 0.62
    avail = X1 - X0 - GAP * (len(PROFILES) - 1)
    if complete:
        coloc = {p: float(data[p]["colocaciones"]) for p in PROFILES}
        tot = sum(coloc.values())
        w = {p: coloc[p] / tot * avail for p in PROFILES}
        floored = [p for p in PROFILES if w[p] < MIN_W]
        if floored:
            fixed = MIN_W * len(floored)
            rest = [p for p in PROFILES if p not in floored]
            rest_tot = sum(coloc[p] for p in rest)
            w = {p: (MIN_W if p in floored else coloc[p] / rest_tot * (avail - fixed)) for p in PROFILES}
    else:
        floored = []
        w = {p: avail / len(PROFILES) for p in PROFILES}

    # alturas = monto / colocaciones (%)
    if complete:
        rate = {p: dict(
            ifn=data[p]["ifn"] / data[p]["colocaciones"] * 100,
            isn=data[p]["isn"] / data[p]["colocaciones"] * 100,
            cr=-abs(data[p]["costo_riesgo"]) / data[p]["colocaciones"] * 100,
        ) for p in PROFILES}
    else:  # silueta de referencia, sin números
        rate = {p: dict(ifn=10.0, isn=4.5, cr=-5.0) for p in PROFILES}
    pos_max = max(max(r["ifn"], r["isn"]) for r in rate.values())
    neg_max = max(-r["cr"] for r in rate.values())
    # margen para las etiquetas: las de barras angostas van en vertical y necesitan más alto
    narrow_any = complete and any((w[p] - 2 * 0.05 - 2 * 0.03) / 3 < 0.30 for p in PROFILES)
    head_t = head_b = 0.37 if narrow_any else 0.17
    k = (plot_bot - plot_top - head_t - head_b) / (pos_max + neg_max)  # pulgadas por punto %
    y0 = plot_top + head_t + pos_max * k

    # ---- fondos de grupo (capa inferior)
    xs, x = {}, X0
    for p in PROFILES:
        xs[p] = x
        add_rect(slide, x, plot_top, w[p], plot_bot - plot_top, fill=BAND)
        x += w[p] + GAP

    # ---- eje Y y grilla
    step = nice_step(pos_max + neg_max)
    v = -math.floor(neg_max / step) * step
    while v <= pos_max + 1e-9:
        y = y0 - v * k
        if complete:
            add_rect(slide, X0 - 0.04, y - 0.003, X1 - X0 + 0.04, 0.006, fill=INK if v == 0 else GRID)
            add_text(slide, 0.41, y - 0.07, 0.56, 0.14, f"{v:g}%", size=7, color=MUTED, align="r")
        v += step
    if not complete:
        add_rect(slide, X0 - 0.04, y0 - 0.003, X1 - X0 + 0.04, 0.006, fill=INK)

    # ---- grupos
    for p in PROFILES:
        gw, d, st, x = w[p], data[p], style[p], xs[p]
        # competidores (arriba)
        comp = d.get("competidores") or {}
        lines = []
        for b in banks:
            val = comp.get(b)
            lines.append([(f"{b} ", {"bold": True, "color": MUTED}),
                          ("pend." if val is None and not complete else "n/d" if val is None else fmt_pct(val, 0 if float(val).is_integer() else 1), {"color": INK if val is not None else MUTED})])
        add_text(slide, x, strip_y, gw, strip_h, lines, size=7, align="c", anchor="m")
        # barras
        pad, ig = 0.05, 0.03
        bw = (gw - 2 * pad - 2 * ig) / 3
        series = (("ifn", ACCENT), ("isn", ACCENT_LIGHT), ("cr", NEG))
        for j, (key, col) in enumerate(series):
            bx = x + pad + j * (bw + ig)
            r = rate[p][key]
            h = abs(r) * k
            by = y0 - h if r >= 0 else y0
            if complete:
                add_rect(slide, bx, by, bw, h, fill=col)
                lab = fmt_pct(r, 1)
                narrow = bw < 0.30
                if narrow:
                    # etiqueta vertical, fuera de la barra
                    ly = by - 0.34 if r >= 0 else by + h + 0.02
                    add_text(slide, bx, ly, bw, 0.32, lab, size=6, color=INK, align="c" if r < 0 else "c",
                             anchor="b" if r >= 0 else "t", vert="vert270")
                else:
                    ly = by - 0.14 if r >= 0 else by + h + 0.01
                    add_text(slide, bx - 0.04, ly, bw + 0.08, 0.13, lab, size=6.5, color=INK, align="c")
            else:
                add_rect(slide, bx, by, bw, h, line="9CA3AF", dash=True)
        # chip del perfil + % de colocaciones
        add_rect(slide, x, chip_y, gw, 0.2, fill=st["fill"], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        full = f"{p} · {st['name']}"
        label = full if len(full) * 0.062 + 0.16 <= gw else p  # el nombre solo si cabe en el chip
        add_text(slide, x, chip_y, gw, 0.2, label, size=7, bold=True, color=st["ink"], align="c")
        if complete:
            share = coloc[p] / tot * 100
            add_text(slide, x, chip_y + 0.22, gw, 0.13, f"{fmt_pct(share, 1)} coloc.", size=6.5, color=MUTED, align="c")
        else:
            add_text(slide, x, chip_y + 0.22, gw, 0.13, "pend.", size=6.5, color=MUTED, align="c")

    # etiqueta de la fila de competidores
    add_text(slide, 0.41, strip_y, 0.60, strip_h, ["Bancos", "competencia", "(% coloc.)"], size=6, color=MUTED, align="r")

    # ---- leyenda
    lx = X0
    for name, col in (("Ingreso financiero neto", ACCENT), ("Ingreso por servicio neto", ACCENT_LIGHT), ("Costo de riesgo", NEG)):
        add_rect(slide, lx, leg_y + 0.025, 0.1, 0.1, fill=col)
        add_text(slide, lx + 0.15, leg_y, 1.7, 0.15, name, size=7.5, color=INK)
        lx += 1.85

    # ---- nota al pie
    nota = ("Ancho de cada grupo: proporcional a las colocaciones del perfil. Alto de cada barra: monto como % de las "
            "colocaciones del perfil, por lo que su área equivale al monto en $. Sobre cada grupo: % de colocaciones de los "
            "bancos competidores.")
    if floored:
        nota += f" Ancho mínimo aplicado a {', '.join(floored)} para que se lea."
    if not complete:
        nota += " Datos pendientes de completar."
    add_text(slide, 0.41, note_y, 9.15, 0.36, nota + " Fuente: pend.", size=7, color=MUTED, anchor="t")

    if demo:
        stamp(slide)

    # mover la lámina nueva para que quede justo después de la 3
    lst = prs.slides._sldIdLst
    el = lst[-1]
    lst.remove(el)
    lst.insert(3, el)
    return complete


# --------------------------------------------------------------------------- main
def main():
    here = Path(__file__).parent
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(here / "Perfiles_Explicados_1.pptx"))
    ap.add_argument("--data", default=str(here / "datos_perfiles.json"))
    ap.add_argument("--out", default=str(here / "Perfiles_Explicados_v2.pptx"))
    ap.add_argument("--demo", action="store_true", help="datos ILUSTRATIVOS para ver el diseño")
    a = ap.parse_args()

    data = load_data(a.data, a.demo)
    prs = Presentation(a.src)
    pcts = pct_clientes(data)
    style = profile_style(prs.slides[2])
    update_profile_table(prs, pcts, a.demo)
    complete = build_sam_slide(prs, data, style, a.demo)
    prs.save(a.out)
    print("OK", a.out)
    print("  % clientes Tenpo:", pcts if pcts["P1"] is not None else "pendiente (falta clientes_tenpo)")
    print("  gráfico SAM:", "completo" if complete else "pendiente (faltan colocaciones/ifn/isn/costo_riesgo)")


if __name__ == "__main__":
    main()
